"""นำเข้าแผนการจัดกิจกรรมทั้งภาคเรียนจากไฟล์ Excel/CSV (ข้อ 6.7).

เจ้าหน้าที่วางแผนกิจกรรมทั้งเทอมในตารางอยู่แล้ว การให้พิมพ์ซ้ำทีละกิจกรรมใน
ฟอร์มเว็บคือการทำงานสองรอบ โมดูลนี้จึงรับไฟล์แผนนั้นตรง ๆ

หลักการที่ยึดไว้: **ไฟล์หนึ่งไฟล์ต้องไม่ล้มทั้งไฟล์เพราะแถวเดียวผิด** — ทุกแถวถูก
ตรวจแยกกัน แถวที่ผ่านถูกสร้างจริง แถวที่ผิดถูกรายงานกลับพร้อมเลขแถวและเหตุผล
ให้ผู้ใช้ไปแก้เฉพาะแถวนั้นแล้วอัปโหลดใหม่

ฟังก์ชันในไฟล์นี้ไม่แตะฐานข้อมูลและไม่แตะ FastAPI (รับ ``bytes`` คืน dataclass)
จึงเทสต์ตรง ๆ ได้ทั้งหมด — ส่วนที่ต้องใช้ DB มีแค่ "ดัชนีหมวดย่อย" ที่ router
สร้างแล้วส่งเข้ามา ``openpyxl`` ถูก import แบบ lazy ตามแนวเดียวกับ reports.py
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any, Optional

# ---------------------------------------------------------------- คอลัมน์

# ชื่อคอลัมน์ที่รับได้ → ชื่อฟิลด์ภายใน (เทียบแบบตัดช่องว่างและไม่สนตัวพิมพ์)
# รับหลายชื่อเพราะไฟล์จริงที่เจ้าหน้าที่ทำไว้ก่อนหน้าอาจใช้หัวตารางไม่ตรงกับ
# template ของเราเป๊ะ ๆ
HEADER_ALIASES: dict[str, str] = {
    "ชื่อ": "name",
    "ชื่อกิจกรรม": "name",
    "name": "name",
    "ประเภท": "activity_type",
    "ประเภทกิจกรรม": "activity_type",
    "activity_type": "activity_type",
    "หมวดย่อย": "subcategory",
    "หมวดชั่วโมง": "subcategory",
    "subcategory": "subcategory",
    "วันเวลา": "start_at",
    "วันที่": "start_at",
    "วันเวลาที่จัด": "start_at",
    "start_at": "start_at",
    "ชั่วโมง": "hours",
    "จำนวนชั่วโมง": "hours",
    "hours": "hours",
    "สถานที่": "location",
    "location": "location",
    "จำนวนรับ": "max_participants",
    "จำนวนที่รับ": "max_participants",
    "max_participants": "max_participants",
    "บังคับ": "is_required",
    "กิจกรรมบังคับ": "is_required",
    "is_required": "is_required",
}

# หัวตารางของ template (เรียงตามสเปกข้อ 6.7)
TEMPLATE_HEADERS = [
    "ชื่อ",
    "ประเภท",
    "หมวดย่อย",
    "วันเวลา",
    "ชั่วโมง",
    "สถานที่",
    "จำนวนรับ",
    "บังคับ",
]

# คอลัมน์ที่ขาดไม่ได้ — "บังคับ" ไม่อยู่ในนี้เพราะเว้นว่างได้ (ถือว่าไม่บังคับ)
REQUIRED_FIELDS = [
    "name",
    "activity_type",
    "subcategory",
    "start_at",
    "hours",
    "location",
    "max_participants",
]

FIELD_LABELS = {
    "name": "ชื่อ",
    "activity_type": "ประเภท",
    "subcategory": "หมวดย่อย",
    "start_at": "วันเวลา",
    "hours": "ชั่วโมง",
    "location": "สถานที่",
    "max_participants": "จำนวนรับ",
    "is_required": "บังคับ",
}

# กันไฟล์ที่ใหญ่เกินกว่าจะเป็น "แผนหนึ่งภาคเรียน" จริง ๆ (กันทั้งไฟล์หลุด
# และกันการยิงงานหนักใส่เซิร์ฟเวอร์ด้วยไฟล์แถวเป็นแสน)
MAX_ROWS = 500

TRUE_WORDS = {"ใช่", "บังคับ", "จริง", "y", "yes", "true", "t", "1"}
FALSE_WORDS = {"ไม่", "ไม่ใช่", "ไม่บังคับ", "เท็จ", "n", "no", "false", "f", "0", "-"}

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEMPLATE_FILENAME = "เทมเพลตนำเข้าแผนกิจกรรม.xlsx"


class ImportFileError(ValueError):
    """ไฟล์ทั้งไฟล์ใช้ไม่ได้ (นามสกุลผิด/ไม่มีหัวตาราง/ว่างเปล่า)

    ต่างจาก "แถวผิด" ตรงที่กรณีนี้ไม่มีอะไรให้ import เลย จึงตอบ 400 ทั้งคำขอ
    แทนการรายงานรายแถว
    """


@dataclass(frozen=True)
class ParsedRow:
    """หนึ่งแถวดิบจากไฟล์ พร้อมเลขแถวจริงในไฟล์ (ไว้บอกผู้ใช้ว่าไปแก้บรรทัดไหน)"""

    row_number: int
    values: dict[str, Any]


@dataclass
class ActivityDraft:
    """แถวที่ผ่านการตรวจแล้ว พร้อมสร้างเป็น Activity"""

    row_number: int
    name: str
    activity_type: str
    subcategory_id: int
    start_at: datetime
    hours: float
    location: str
    max_participants: int
    is_required: bool = False


@dataclass
class RowIssue:
    row_number: int
    message: str


@dataclass
class ValidationOutcome:
    drafts: list[ActivityDraft] = field(default_factory=list)
    issues: list[RowIssue] = field(default_factory=list)


# ---------------------------------------------------------------- อ่านไฟล์


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def _map_headers(cells: list[Any]) -> dict[int, str]:
    """คอลัมน์ที่ index ไหน = ฟิลด์อะไร (คอลัมน์ที่ไม่รู้จักถูกเมินเฉย ๆ)"""
    mapping: dict[int, str] = {}
    for index, cell in enumerate(cells):
        field_name = HEADER_ALIASES.get(_normalize_header(cell))
        if field_name and field_name not in mapping.values():
            mapping[index] = field_name
    return mapping


def _require_all_headers(mapping: dict[int, str]) -> None:
    missing = [FIELD_LABELS[f] for f in REQUIRED_FIELDS if f not in mapping.values()]
    if missing:
        raise ImportFileError(
            "ไฟล์ขาดคอลัมน์: " + ", ".join(missing) + " (ดาวน์โหลดไฟล์ต้นแบบแล้วกรอกตามหัวตาราง)"
        )


def _rows_from_grid(grid: list[tuple[int, list[Any]]]) -> list[ParsedRow]:
    """หา "แถวหัวตาราง" แล้วอ่านแถวข้อมูลที่เหลือ

    ``grid`` คือ (เลขแถวจริงในไฟล์, ค่าในแถว) — ไฟล์ที่คนทำเองมักมีบรรทัดชื่อเรื่อง
    อยู่ข้างบนก่อนหัวตาราง จึงไล่หาแถวแรกที่หน้าตาเป็นหัวตารางแทนการฟันธงว่าแถว 1
    """
    header_index: Optional[int] = None
    mapping: dict[int, str] = {}
    for position, (_, cells) in enumerate(grid):
        candidate = _map_headers(cells)
        if len(set(candidate.values()) & set(REQUIRED_FIELDS)) >= 3:
            header_index = position
            mapping = candidate
            break

    if header_index is None:
        raise ImportFileError("ไม่พบหัวตารางในไฟล์ (ดาวน์โหลดไฟล์ต้นแบบแล้วกรอกตามหัวตาราง)")
    _require_all_headers(mapping)

    rows: list[ParsedRow] = []
    for row_number, cells in grid[header_index + 1 :]:
        values = {
            field_name: (cells[index] if index < len(cells) else None)
            for index, field_name in mapping.items()
        }
        # ข้ามแถวที่ว่างทั้งแถว (Excel แถมแถวว่างท้ายไฟล์บ่อยมาก)
        if all(_is_blank(v) for v in values.values()):
            continue
        rows.append(ParsedRow(row_number=row_number, values=values))

    if not rows:
        raise ImportFileError("ไฟล์ไม่มีข้อมูลกิจกรรมสักแถว")
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"ไฟล์มีข้อมูล {len(rows)} แถว เกินที่รับได้ {MAX_ROWS} แถวต่อครั้ง")
    return rows


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _decode_csv(content: bytes) -> str:
    """ถอดรหัส CSV: UTF-8 ก่อน แล้วค่อย cp874

    Excel ภาษาไทยบน Windows บันทึก CSV เป็น cp874 (TIS-620) ไม่ใช่ UTF-8 —
    ถ้าไม่รองรับ ผู้ใช้จะเจอชื่อกิจกรรมเป็นอักขระเพี้ยนทั้งไฟล์โดยไม่รู้สาเหตุ
    """
    for encoding in ("utf-8-sig", "cp874"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ImportFileError("อ่านไฟล์ CSV ไม่ได้ (ต้องบันทึกเป็น UTF-8 หรือ TIS-620)")


def parse_spreadsheet(filename: str, content: bytes) -> list[ParsedRow]:
    """อ่านไฟล์ .xlsx/.csv เป็นแถวดิบ — ยังไม่ตรวจความถูกต้องของค่า"""
    if not content:
        raise ImportFileError("ไฟล์ว่างเปล่า")

    lowered = (filename or "").lower()
    if lowered.endswith(".xlsx"):
        grid = _grid_from_xlsx(content)
    elif lowered.endswith(".csv"):
        text = _decode_csv(content)
        reader = csv.reader(io.StringIO(text))
        grid = [(number, list(cells)) for number, cells in enumerate(reader, start=1)]
    else:
        raise ImportFileError("รองรับเฉพาะไฟล์ .xlsx และ .csv")

    if not grid:
        raise ImportFileError("ไฟล์ว่างเปล่า")
    return _rows_from_grid(grid)


def _grid_from_xlsx(content: bytes) -> list[tuple[int, list[Any]]]:
    from openpyxl import load_workbook  # lazy — ดู docstring ของโมดูล

    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl โยน exception ได้หลายชนิด
        raise ImportFileError("เปิดไฟล์ Excel ไม่ได้ (ไฟล์อาจเสียหรือไม่ใช่ .xlsx จริง)") from exc

    sheet = workbook.active
    return [(number, list(cells)) for number, cells in enumerate(sheet.iter_rows(values_only=True), start=1)]


# ---------------------------------------------------------------- ตรวจค่า


def normalize_lookup_key(value: Any) -> str:
    """คีย์สำหรับเทียบชื่อหมวด: ตัดช่องว่างทั้งหมดและไม่สนตัวพิมพ์

    ไฟล์ที่คนกรอกเองมีช่องว่างเกินท้ายชื่อเสมอ และ "ICT" กับ "ict" ต้องเป็นอันเดียวกัน
    """
    return re.sub(r"\s+", "", str(value or "")).lower()


def _parse_datetime(value: Any) -> datetime:
    """แปลงค่าในช่อง "วันเวลา" เป็น datetime

    รับทั้งเซลล์วันเวลาจริงของ Excel และข้อความหลายรูปแบบ รวมถึง **ปี พ.ศ.**
    (2569 → 2026) เพราะไฟล์ที่เจ้าหน้าที่ไทยพิมพ์เองมักเป็น พ.ศ. ถ้าไม่แปลงให้
    จะกลายเป็นกิจกรรมปี ค.ศ. 2569 ซึ่งผ่านกฎ "ห้ามย้อนหลัง" ไปแบบเงียบ ๆ
    """
    if isinstance(value, datetime):
        return _to_christian_era(value)
    if isinstance(value, date):
        return _to_christian_era(datetime.combine(value, time.min))

    text = str(value or "").strip()
    if not text:
        raise ValueError("ไม่ได้กรอกวันเวลา")
    text = text.replace("T", " ")
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
                    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return _to_christian_era(datetime.strptime(text, pattern))
        except ValueError:
            continue
    raise ValueError(f'วันเวลา "{text}" อ่านไม่ออก (ใช้รูปแบบ 2026-12-01 09:00)')


def _to_christian_era(value: datetime) -> datetime:
    # ปีเกิน 2400 เป็นไปไม่ได้ในเชิงปฏิทิน ค.ศ. สำหรับกิจกรรมมหาวิทยาลัย → เป็น พ.ศ.
    return value.replace(year=value.year - 543) if value.year >= 2400 else value


def _parse_positive_number(value: Any, label: str) -> float:
    text = str(value or "").strip()
    if not text:
        raise ValueError(f"ไม่ได้กรอก{label}")
    try:
        number = float(text)
    except ValueError:
        raise ValueError(f'{label} "{text}" ไม่ใช่ตัวเลข') from None
    if number <= 0:
        raise ValueError(f"{label}ต้องมากกว่า 0")
    return number


def _parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    text = str(value or "").strip().lower()
    if not text:
        return False  # เว้นว่าง = ไม่บังคับ
    if text in TRUE_WORDS:
        return True
    if text in FALSE_WORDS:
        return False
    raise ValueError(f'ช่องบังคับ "{text}" อ่านไม่ออก (ใช้ ใช่/ไม่ใช่)')


def _parse_text(value: Any, label: str) -> str:
    text = str(value or "").strip()
    if not text:
        raise ValueError(f"ไม่ได้กรอก{label}")
    return text


def validate_rows(
    rows: list[ParsedRow],
    *,
    subcategory_index: dict[str, list[int]],
    today: date,
    existing_keys: Optional[set[tuple[str, datetime]]] = None,
) -> ValidationOutcome:
    """ตรวจทุกแถว คืนทั้งแถวที่ผ่านและเหตุผลของแถวที่ไม่ผ่าน

    ``subcategory_index`` map จากคีย์ที่ผู้ใช้กรอกได้ (ชื่อหมวดย่อย / "หมวดแม่ ›
    หมวดย่อย" / เลข id) ไปยัง id ที่ตรงกัน — เป็น list เพราะชื่อหมวดย่อยซ้ำข้าม
    หมวดแม่ได้ ถ้าซ้ำจะให้ผู้ใช้ระบุเป็นชื่อเต็มแทนการเดาให้

    ``existing_keys`` คือ (ชื่อ, วันเวลา) ของกิจกรรมที่มีอยู่แล้ว ใช้กันการอัปโหลด
    ไฟล์เดิมซ้ำสองรอบแล้วได้กิจกรรมซ้ำทั้งแผน (แถวที่ซ้ำกันเองในไฟล์ก็ถูกกันด้วย)
    """
    outcome = ValidationOutcome()
    seen: set[tuple[str, datetime]] = set(existing_keys or set())

    for row in rows:
        errors: list[str] = []
        values = row.values

        try:
            name = _parse_text(values.get("name"), "ชื่อกิจกรรม")
        except ValueError as exc:
            name = ""
            errors.append(str(exc))
        try:
            activity_type = _parse_text(values.get("activity_type"), "ประเภท")
        except ValueError as exc:
            activity_type = ""
            errors.append(str(exc))
        try:
            location = _parse_text(values.get("location"), "สถานที่")
        except ValueError as exc:
            location = ""
            errors.append(str(exc))

        start_at: Optional[datetime] = None
        try:
            start_at = _parse_datetime(values.get("start_at"))
            if start_at.date() < today:
                errors.append("วันเวลากิจกรรมเป็นอดีต")
        except ValueError as exc:
            errors.append(str(exc))

        hours = 0.0
        try:
            hours = _parse_positive_number(values.get("hours"), "ชั่วโมง")
        except ValueError as exc:
            errors.append(str(exc))

        max_participants = 0
        try:
            max_participants = int(_parse_positive_number(values.get("max_participants"), "จำนวนรับ"))
        except ValueError as exc:
            errors.append(str(exc))

        is_required = False
        try:
            is_required = _parse_bool(values.get("is_required"))
        except ValueError as exc:
            errors.append(str(exc))

        subcategory_id: Optional[int] = None
        raw_subcategory = str(values.get("subcategory") or "").strip()
        if not raw_subcategory:
            errors.append("ไม่ได้กรอกหมวดย่อย")
        else:
            matches = subcategory_index.get(normalize_lookup_key(raw_subcategory), [])
            if not matches:
                errors.append(f'ไม่พบหมวดย่อย "{raw_subcategory}" ในระบบ')
            elif len(matches) > 1:
                errors.append(
                    f'หมวดย่อย "{raw_subcategory}" มีมากกว่าหนึ่งหมวด '
                    "ให้ระบุเป็น หมวดแม่ › หมวดย่อย"
                )
            else:
                subcategory_id = matches[0]

        if not errors and start_at is not None:
            key = (normalize_lookup_key(name), start_at)
            if key in seen:
                errors.append("ซ้ำกับกิจกรรมชื่อเดียวกันในวันเวลาเดียวกัน")
            else:
                seen.add(key)

        if errors:
            outcome.issues.append(RowIssue(row_number=row.row_number, message=" · ".join(errors)))
            continue

        assert start_at is not None and subcategory_id is not None  # ผ่านการตรวจข้างบนแล้ว
        outcome.drafts.append(
            ActivityDraft(
                row_number=row.row_number,
                name=name,
                activity_type=activity_type,
                subcategory_id=subcategory_id,
                start_at=start_at,
                hours=hours,
                location=location,
                max_participants=max_participants,
                is_required=is_required,
            )
        )

    return outcome


# ---------------------------------------------------------------- ไฟล์ต้นแบบ


def build_template_xlsx(subcategory_paths: list[str]) -> bytes:
    """ไฟล์ต้นแบบเปล่า: หัวตาราง + ตัวอย่าง 1 แถว + ชีตรายชื่อหมวดย่อยที่ใช้ได้จริง

    ชีตหมวดย่อยสร้างสดจากฐานข้อมูล เพราะสาเหตุที่แถวไม่ผ่านบ่อยที่สุดคือพิมพ์ชื่อ
    หมวดย่อยไม่ตรง — ให้ก็อปจากไฟล์ต้นแบบไปวางได้เลยจะเร็วกว่าไล่เดา
    """
    from openpyxl import Workbook  # lazy
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "แผนกิจกรรม"

    sheet.append(TEMPLATE_HEADERS)
    header_fill = PatternFill("solid", start_color="FFE8EAF6")
    for column in range(1, len(TEMPLATE_HEADERS) + 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    example_subcategory = subcategory_paths[0] if subcategory_paths else "หมวดแม่ › หมวดย่อย"
    sheet.append(
        [
            "ตัวอย่าง: อบรมการใช้ห้องสมุด",
            "อบรม",
            example_subcategory,
            "2026-12-01 09:00",
            3,
            "ห้องประชุม 1",
            80,
            "ไม่ใช่",
        ]
    )

    sheet.freeze_panes = "A2"
    for index, width in enumerate([34, 16, 40, 20, 10, 24, 12, 10], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    notes = workbook.create_sheet("หมวดย่อยที่ใช้ได้")
    notes.append(["คัดลอกชื่อจากคอลัมน์นี้ไปวางในช่อง หมวดย่อย"])
    notes["A1"].font = Font(bold=True)
    for path in subcategory_paths:
        notes.append([path])
    notes.column_dimensions["A"].width = 60

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
