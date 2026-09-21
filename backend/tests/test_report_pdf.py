"""รายงาน PDF: เลย์เอาต์ตาราง, รูปแบบ "สรุปต่อคน", ตัวกรองภาคเรียน, ตัวนับแถว

PDF ถูกเปิดอ่านกลับจริงด้วย PyMuPDF (ข้อความ, ตำแหน่งคำ, ฟอนต์ที่ฝัง, เลขหน้า) — เทสต์
เลย์เอาต์ที่ดูแค่ "ได้ไฟล์ %PDF" จะไม่จับบั๊กแบบคณะล้นไปชนชั้นปีที่เกิดขึ้นจริง
"""

import pymupdf
import pytest
from reportlab.pdfbase.pdfmetrics import stringWidth

from app.fonts import find_report_fonts
from app.report_pdf import (
    _FIT_SLACK,
    PdfColumn,
    build_pdf,
    build_summary_pdf,
    column_widths,
    detail_columns,
)
from app.reports import (
    LARGE_REPORT_ROWS,
    STATUS_COMPLETE,
    STATUS_INCOMPLETE,
    ReportRow,
    StudentSummary,
    count_report_rows,
    describe_filters,
    estimate_pdf_pages,
    fetch_report_rows,
    fetch_report_summary,
    summarize_rows,
)
from app.thai_text import has_stacked_marks, stacked_marks_markup, wrap_text
from tests.test_reports import _admin_headers, _student_headers, world_fixture  # noqa: F401

# ชื่อคณะ/หน่วยงานจริงทั้ง 11 ที่มีในไฟล์นำเข้านิสิต (ยาวสุดคือ "คณะวิทยาศาสตร์และนวัตกรรมดิจิทัล")
REAL_FACULTIES = [
    "คณะวิทยาการสุขภาพและการกีฬา",
    "คณะวิศวกรรมศาสตร์",
    "คณะศึกษาศาสตร์",
    "คณะนิติศาสตร์",
    "วิทยาลัยการจัดการเพื่อการพัฒนา",
    "คณะวิทยาศาสตร์และนวัตกรรมดิจิทัล",
    "คณะสหเวชศาสตร์",
    "คณะเกษตรศาสตร์",
    "คณะพยาบาลศาสตร์",
    "คณะอุตสาหกรรมเกษตรและชีวภาพ",
    "คณะสหวิทยาการและการประกอบการ",
]
CATEGORIES = [
    "พัฒนาทักษะชีวิต",
    "TSU รับใช้สังคม",
    "ศิลปวัฒนธรรมอาเซียน",
    "พัฒนาคุณธรรมและวินัย",
    "ใฝ่เรียนรู้ตลอดชีวิต",
]


def _font():
    from app.report_pdf import _register_thai_font

    return _register_thai_font()


def _rows(n_students, faculty="คณะวิทยาศาสตร์และนวัตกรรมดิจิทัล", name="นางสาวสมชาย ใจดีมีสุข"):
    return [
        ReportRow(
            f"68{i:08d}", name, faculty, 1 + i % 4, cat, 12 if c % 2 else 0, 12, bool(c % 2), student_key=i
        )
        for i in range(n_students)
        for c, cat in enumerate(CATEGORIES)
    ]


def _open(content: bytes):
    return pymupdf.open(stream=content, filetype="pdf")


# ---------------------------------------------------------------- ตัดบรรทัด / วรรณยุกต์
def _measure():
    font, _ = _font()
    return lambda v: stringWidth(v, font, 9)


@pytest.mark.parametrize("width", [40, 60, 90, 140])
@pytest.mark.parametrize("text", REAL_FACULTIES + CATEGORIES + ["นายสมปองสมหมายปรารถนาดีมีชัยวัฒนาศิริมงคลบุญเจริญ ผู้เข้าร่วม"])
def test_wrap_never_exceeds_width_or_loses_text(text, width):
    measure = _measure()
    lines = wrap_text(text, width, measure)
    assert "".join(lines).replace(" ", "") == text.replace(" ", "")  # ไม่หายไม่เพิ่มตัวอักษร
    for line in lines:
        assert measure(line) <= width + 0.01, (line, width)


def test_wrap_leaves_short_text_untouched_and_handles_empty():
    measure = _measure()
    assert wrap_text("คณะนิติศาสตร์", 500, measure) == ["คณะนิติศาสตร์"]
    assert wrap_text("", 100, measure) == [""]


def test_wrap_never_starts_a_line_with_a_mark_or_strands_a_leading_vowel():
    import unicodedata

    measure = _measure()
    for text in REAL_FACULTIES + ["เพื่อการเรียนรู้ตลอดชีวิตและการพัฒนาที่ยั่งยืนของแม่โจ้"]:
        for width in range(25, 120, 7):
            lines = wrap_text(text, width, measure)
            for line in lines:
                assert unicodedata.category(line[0]) != "Mn" and line[0] not in "ะาำๆฯ", (lines, width)
                assert line.rstrip()[-1] not in "เแโใไ", (lines, width)


def test_wrap_prefers_breaking_at_a_conjunction():
    measure = _measure()
    lines = wrap_text("คณะวิทยาศาสตร์และนวัตกรรมดิจิทัล", measure("คณะวิทยาศาสตร์") + 4, measure)
    assert lines[0] == "คณะวิทยาศาสตร์"
    assert lines[1].startswith("และ")


def test_stacked_tone_marks_are_raised_and_plain_text_is_not():
    for word in ["ที่", "นี้", "ซึ่ง", "ชั่ว", "ป่า", "ปู่", "ฝ่าย", "ฟ้า", "ก๊ก"]:
        marked = stacked_marks_markup(word, 9)
        # ก๊ก = พยัญชนะ+วรรณยุกต์ ไม่ซ้อนอะไรเลย ไม่ต้องยก
        assert ("<super" in marked) == (word != "ก๊ก"), word
        assert has_stacked_marks(word) == (word != "ก๊ก")
    assert stacked_marks_markup("A & B <x>", 9) == "A &amp; B &lt;x&gt;"  # escape ให้ Paragraph ไม่พัง


# ---------------------------------------------------------------- ความกว้างคอลัมน์
def _widths_for(faculty_values):
    font, bold = _font()
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm

    rows = [
        ReportRow("6800000001", "นางสาวสมชาย ใจดี", f, 2, "พัฒนาทักษะชีวิต", 1, 12, False)
        for f in faculty_values
    ]
    columns = detail_columns(rows)
    available = landscape(A4)[0] - 24 * mm
    pad = 2.5 * mm
    measure = lambda v: max(stringWidth(v, font, 9), stringWidth(v, bold, 9))  # noqa: E731
    return columns, column_widths(columns, available, measure, pad), available, pad, measure


def test_faculty_column_fits_the_longest_real_faculty_on_one_line():
    """ต้นเหตุของบั๊กเดิม: คอลัมน์ "คณะ" แคบกว่าชื่อคณะที่ยาวที่สุด จึงล้นไปชนชั้นปี"""
    columns, widths, available, pad, _measure_fn = _widths_for(REAL_FACULTIES)
    font, _ = _font()
    faculty_index = 2
    for faculty in REAL_FACULTIES:
        need = stringWidth(faculty, font, 9) + 2 * pad
        assert widths[faculty_index] >= need, (faculty, widths[faculty_index], need)
    assert sum(widths) == pytest.approx(available)


def test_widths_never_exceed_the_page_and_shrink_only_text_columns():
    font, _ = _font()
    long_text = "ก" * 400
    columns = [
        PdfColumn("รหัส", ["6800000001"], "LEFT"),
        PdfColumn("ชื่อ", [long_text], "LEFT", flex=True),
        PdfColumn("คณะ", [long_text], "LEFT", flex=True),
        PdfColumn("ปี", ["4"], "CENTER"),
    ]
    measure = lambda v: stringWidth(v, font, 9)  # noqa: E731
    widths = column_widths(columns, 500, measure, 7)
    assert sum(widths) <= 500.5
    assert widths[0] >= measure("6800000001") + 14 + _FIT_SLACK - 0.01  # คอลัมน์รหัส/ปีไม่ถูกบีบ
    assert widths[3] >= measure("ปี") + 14 - 0.01


# ---------------------------------------------------------------- ไฟล์ PDF
def test_bundled_sarabun_is_used_and_embedded():
    regular, bold = find_report_fonts()
    assert regular.endswith("Sarabun-Regular.ttf") and bold.endswith("Sarabun-Bold.ttf")
    content = build_pdf(_rows(3), filter_note="ทุกคณะ")
    assert b"FontFile2" in content  # TrueType ฝังมาในไฟล์ ไม่ใช่แค่อ้างชื่อ
    names = {f[3] for f in _open(content)[0].get_fonts()}
    # ฟอนต์ที่ฝังจริงเป็น subset ของ Sarabun ทั้งตัวปกติและตัวหนา (Helvetica คือค่าเริ่มต้นของ
    # reportlab ที่ไม่ได้ใช้วาดตัวอักษรไทย)
    assert any(n.endswith("Sarabun-Regular") for n in names), names
    assert any(n.endswith("Sarabun-Bold") for n in names), names


def test_thai_text_survives_round_trip_including_stacked_marks():
    """ถ้าวรรณยุกต์หลุด/ฟอนต์ไม่ครบ ข้อความที่อ่านกลับจากไฟล์จะไม่ตรงต้นฉบับ"""
    doc = _open(build_pdf(_rows(2), filter_note="คณะวิทยาศาสตร์และนวัตกรรมดิจิทัล"))
    text = doc[0].get_text()
    for expected in [
        "รายงานชั่วโมงกิจกรรมสะสมของนิสิต",
        "ชั่วโมงที่ได้",
        "ชั่วโมงที่ต้องการ",
        "คณะวิทยาศาสตร์และนวัตกรรมดิจิทัล",
        "พัฒนาคุณธรรมและวินัย",
        "ใฝ่เรียนรู้ตลอดชีวิต",
        "ไม่ครบ",
    ]:
        assert expected in text, expected
    assert "�" not in text and "\x00" not in text


def test_every_page_has_page_x_of_y_and_repeated_header():
    doc = _open(build_pdf(_rows(40), filter_note="ทุกคณะ"))  # 200 แถว → หลายหน้า
    assert doc.page_count > 3
    for number, page in enumerate(doc, start=1):
        text = page.get_text()
        assert f"หน้า {number}/{doc.page_count}" in text
        assert "รหัสนิสิต" in text and "ชั่วโมงที่ต้องการ" in text  # หัวตารางซ้ำทุกหน้า


def _cell_boxes(page):
    tables = page.find_tables()
    assert tables.tables, "หาตารางในหน้าไม่เจอ"
    return [pymupdf.Rect(c) for t in tables.tables for c in t.cells if c]


def test_no_text_spills_out_of_its_cell_even_with_very_long_text():
    """ข้อความยาวมาก (ไม่มีช่องว่างเลย) ต้องถูกตัดบรรทัดในเซลล์ ห้ามล้นไปชนคอลัมน์ข้าง ๆ"""
    rows = _rows(6, faculty="คณะเกษตรศาสตร์และทรัพยากรธรรมชาติและสิ่งแวดล้อมเพื่อการพัฒนาที่ยั่งยืนแห่งภูมิภาค",
                 name="นายสมปองสมหมายปรารถนาดีมีชัยวัฒนาศิริมงคลบุญเจริญผู้เข้าร่วมกิจกรรมพิเศษ")
    page = _open(build_pdf(rows, filter_note="ทดสอบ"))[0]
    cells = _cell_boxes(page)
    # เผื่อ 0.5pt ให้ขอบเส้น — คำทุกคำต้องอยู่ในเซลล์ใดเซลล์หนึ่งทั้งคำ
    for x0, y0, x1, y1, word, *_ in page.get_text("words"):
        rect = pymupdf.Rect(x0, y0, x1, y1)
        if rect.y0 < cells[0].y0 or word.startswith("หน้า") or "/" in word and word.replace("/", "").isdigit():
            continue  # หัวรายงาน/เลขหน้า อยู่นอกตาราง
        assert any(
            c.x0 - 0.5 <= rect.x0 and rect.x1 <= c.x1 + 0.5 and c.y0 - 0.5 <= rect.y0 and rect.y1 <= c.y1 + 0.5
            for c in cells
        ), f"คำ {word!r} ล้นเซลล์ {rect}"


def test_cells_have_horizontal_padding():
    """ตัวอักษรห้ามติดเส้นขอบเซลล์"""
    page = _open(build_pdf(_rows(4), filter_note="x"))[0]
    cells = _cell_boxes(page)
    table_bottom = max(c.y1 for c in cells)
    words = [
        pymupdf.Rect(w[:4]) for w in page.get_text("words") if cells[0].y0 < w[1] and w[3] <= table_bottom
    ]  # เฉพาะคำในตาราง (ตัดเลขหน้าท้ายกระดาษ)
    assert words
    for rect in words:
        cell = next(c for c in cells if c.contains(rect + (-0.5, -0.5, 0.5, 0.5)))
        left_gap, right_gap = rect.x0 - cell.x0, cell.x1 - rect.x1
        assert min(left_gap, right_gap) >= 4, (rect, cell)  # ≥ 4pt ทั้งสองข้าง (padding ตั้งไว้ ~7pt)


def test_year_level_column_is_centred():
    page = _open(build_pdf(_rows(1, faculty="คณะนิติศาสตร์"), filter_note="x"))[0]
    cells = _cell_boxes(page)
    header = next(w for w in page.get_text("words") if w[4] == "ชั้นปี")
    year_cell = next(c for c in cells if c.x0 <= header[0] and header[2] <= c.x1 and c.y0 > header[3])
    digit = next(w for w in page.get_text("words") if w[4] == "1" and year_cell.contains(pymupdf.Rect(w[:4])))
    assert abs((digit[0] + digit[2]) / 2 - (year_cell.x0 + year_cell.x1) / 2) < 1.5


def test_rows_per_page_estimate_matches_real_page_count():
    students = 120  # 600 แถว
    pages = _open(build_pdf(_rows(students), filter_note="x")).page_count
    estimate = estimate_pdf_pages(students * len(CATEGORIES))
    assert abs(pages - estimate) <= max(2, pages * 0.12), (pages, estimate)


# ---------------------------------------------------------------- สรุปต่อคน
def test_summarize_rows_one_row_per_student_with_counts_and_status():
    rows = [
        ReportRow("A1", "เอ", "คณะ ก", 1, "หมวด1", 10, 10, True, student_key=1),
        ReportRow("A1", "เอ", "คณะ ก", 1, "หมวด2", 3, 5, False, student_key=1),
        ReportRow("B2", "บี", "คณะ ข", 2, "หมวด1", 10, 10, True, student_key=2),
        ReportRow("B2", "บี", "คณะ ข", 2, "หมวด2", 5, 5, True, student_key=2),
    ]
    a, b = summarize_rows(rows, {1: 13.0, 2: 15.0})
    assert (a.student_code, a.total_hours, a.completed_items, a.total_items) == ("A1", 13.0, 1, 2)
    assert a.items_text == "ครบ 1/2 หมวด" and a.status_text == STATUS_INCOMPLETE
    assert b.items_text == "ครบ 2/2 หมวด" and b.status_text == STATUS_COMPLETE


def test_student_with_no_items_is_never_marked_complete():
    assert StudentSummary("X", "เอ็กซ์", "คณะ", 1, 0, 0, 0).status_text == STATUS_INCOMPLETE


def test_summary_pdf_has_one_row_per_student_and_fewer_pages_than_detail():
    rows = _rows(60)
    summaries = summarize_rows(rows, {})
    assert len(summaries) == 60
    summary_doc = _open(build_summary_pdf(summaries, filter_note="x"))
    detail_doc = _open(build_pdf(rows, filter_note="x"))
    assert summary_doc.page_count * 3 < detail_doc.page_count  # 60 แถว vs 300 แถว
    text = "".join(p.get_text() for p in summary_doc)
    assert text.count("ครบ 2/5 หมวด") == 60  # แถวละ 1 ครั้ง (c%2 → ครบ 2 จาก 5)
    assert "นิสิตทั้งหมด 60 คน" in text
    assert "หมวดที่ครบ" in text and "ชั่วโมงที่ต้องการ" not in text


def test_summary_endpoint_row_count_and_total_hours_match_the_data(client, tokens, world, session):
    response = client.get("/reports/student-hours.pdf", headers=_admin_headers(tokens))
    assert response.status_code == 200
    doc = _open(response.content)
    text = doc[0].get_text()
    assert "นิสิตทั้งหมด 2 คน" in text  # ค่าเริ่มต้น = สรุปต่อคน (2 นิสิต ไม่ใช่ 4 แถวรายหมวด)
    assert "ครบ 1/2 หมวด" in text  # 660001 ครบสังคม ยังไม่ได้กีฬา
    assert "ครบ 0/2 หมวด" in text  # 660002 ยังไม่มีชั่วโมง

    summaries = {s.student_code: s for s in fetch_report_summary(session)}
    assert summaries["660001"].total_hours == 10
    assert summaries["660002"].total_hours == 0


def test_detail_mode_is_still_available_and_one_row_per_category(client, tokens, world):
    response = client.get(
        "/reports/student-hours.pdf", params={"mode": "detail"}, headers=_admin_headers(tokens)
    )
    text = _open(response.content)[0].get_text()
    assert "ทั้งหมด 4 รายการ (นิสิต 2 คน)" in text
    assert "ชั่วโมงที่ต้องการ" in text and "กิจกรรมกีฬา" in text


def test_unknown_mode_and_semester_are_rejected(client, tokens, world):
    admin = _admin_headers(tokens)
    assert client.get("/reports/student-hours.pdf", params={"mode": "huge"}, headers=admin).status_code == 422
    assert client.get("/reports/student-hours.pdf", params={"semester": 4}, headers=admin).status_code == 422
    assert client.get("/reports/student-hours.xlsx", params={"semester": 0}, headers=admin).status_code == 422


# ---------------------------------------------------------------- ตัวกรอง
def test_pdf_endpoint_outputs_only_the_filtered_faculty(client, tokens, world):
    admin = _admin_headers(tokens)
    for mode in ("summary", "detail"):
        text = _open(
            client.get(
                "/reports/student-hours.pdf",
                params={"faculty": "วิทยาศาสตร์", "mode": mode},
                headers=admin,
            ).content
        )[0].get_text()
        assert "660002" in text and "660001" not in text, mode
        assert "เงื่อนไข: คณะวิทยาศาสตร์" in text


def test_filter_note_does_not_double_the_faculty_prefix_and_names_the_semester():
    assert describe_filters("คณะนิติศาสตร์", None) == "คณะนิติศาสตร์"
    assert describe_filters("วิทยาลัยการจัดการเพื่อการพัฒนา", 2).startswith("วิทยาลัยการจัดการเพื่อการพัฒนา · ")
    note = describe_filters(None, None, 3)
    assert "ภาคฤดูร้อน" in note and "ทั้งหลักสูตร" in note  # บอกชัดว่าชั่วโมงที่ต้องการไม่ถูกหารตามภาค


def test_semester_filter_recomputes_earned_hours_only_for_that_semester(session, world):
    """กิจกรรมเริ่ม 1 มิ.ย. → ภาคเรียนที่ 1 (มิ.ย.–ต.ค.)"""

    def social(semester):
        rows = fetch_report_rows(session, semester=semester)
        return next(r for r in rows if r.student_code == "660001" and r.category_name == "กิจกรรมเพื่อสังคม")

    everything, first, second = social(None), social(1), social(2)
    assert (everything.earned_hours, everything.completed) == (10, True)
    assert (first.earned_hours, first.completed) == (10, True)
    assert (second.earned_hours, second.completed) == (0, False)  # ไม่มีกิจกรรมในภาค 2
    assert second.required_hours == 10  # ชั่วโมงที่ต้องการยังเป็นของทั้งหลักสูตร

    totals = {s.student_code: s.total_hours for s in fetch_report_summary(session, semester=2)}
    assert totals["660001"] == 0
    totals = {s.student_code: s.total_hours for s in fetch_report_summary(session, semester=1)}
    assert totals["660001"] == 10


def test_semester_filter_is_applied_by_both_endpoints(client, tokens, world):
    admin = _admin_headers(tokens)
    pdf = client.get("/reports/student-hours.pdf", params={"semester": 2}, headers=admin)
    text = _open(pdf.content)[0].get_text()
    assert "ภาคเรียนที่ 2" in text and "ครบ 0/2 หมวด" in text  # ไม่มีใครครบอะไรในภาค 2

    from openpyxl import load_workbook
    import io

    xlsx = client.get("/reports/student-hours.xlsx", params={"semester": 2}, headers=admin)
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    rows = [[c.value for c in r] for r in sheet.iter_rows()]
    data = [r for r in rows if r[0] == "660001"]
    assert data and all(r[5] == 0 for r in data)  # ชั่วโมงที่ได้ของภาค 2 เป็น 0 ทุกหมวด


# ---------------------------------------------------------------- ตัวนับแถว (เตือนก่อนออกไฟล์ใหญ่)
def test_count_endpoint_reports_rows_pages_and_threshold(client, tokens, world):
    body = client.get("/reports/student-hours/count", headers=_admin_headers(tokens)).json()
    assert (body["students"], body["summary_rows"], body["detail_rows"]) == (2, 2, 4)
    assert body["large_threshold"] == LARGE_REPORT_ROWS == 5000
    assert body["summary_pages"] == 1 and body["detail_pages"] == 1

    filtered = client.get(
        "/reports/student-hours/count", params={"faculty": "วิทยาศาสตร์"}, headers=_admin_headers(tokens)
    ).json()
    assert (filtered["students"], filtered["detail_rows"]) == (1, 2)


def test_count_matches_what_the_files_contain(session, world):
    students, detail = count_report_rows(session)
    assert detail == len(fetch_report_rows(session))
    assert students == len(fetch_report_summary(session))


def test_count_endpoint_permissions(client, tokens, world):
    assert client.get("/reports/student-hours/count").status_code == 200  # staff
    assert client.get("/reports/student-hours/count", headers=_student_headers(tokens)).status_code == 403


def test_estimate_pdf_pages_rounds_up_and_is_at_least_one():
    assert estimate_pdf_pages(0) == 1
    assert estimate_pdf_pages(1) == 1
    assert estimate_pdf_pages(24) == 1
    assert estimate_pdf_pages(25) == 2
