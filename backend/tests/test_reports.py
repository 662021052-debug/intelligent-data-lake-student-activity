"""ส่งออกรายงานชั่วโมงเป็น Excel / PDF (ข้อ 6.5) — Prompt 2

ไฟล์ที่ได้ถูกอ่านกลับมาตรวจจริง (openpyxl เปิด workbook, PDF ตรวจ magic bytes)
ไม่ใช่แค่ดูว่า endpoint ตอบ 200 — ไฟล์ส่งออกที่เปิดไม่ได้คือความพังที่เงียบที่สุด
"""

import io
from datetime import datetime

import pytest
from sqlmodel import select

from app.gold import create_gold_layer
from app.models import (
    Activity,
    ApprovalStatus,
    EvidenceStatus,
    HourCategory,
    HourSubcategory,
    Participation,
    Student,
    StudentStatus,
    User,
)
from app.reports import (
    COLUMN_HEADERS,
    STATUS_COMPLETE,
    STATUS_INCOMPLETE,
    build_xlsx,
    describe_filters,
    fetch_report_rows,
    report_filename,
)


def _admin_headers(tokens):
    return {"Authorization": f"Bearer {tokens['admin']}"}


def _student_headers(tokens):
    return {"Authorization": f"Bearer {tokens['student']}"}


def _staff_id(session):
    return session.exec(select(User).where(User.username == "staff")).first().id


@pytest.fixture(name="world")
def world_fixture(session):
    """2 หมวด (10 + 5 ชม.) × นิสิต 2 คนคนละคณะ — พอให้ตรวจทั้งครบ/ไม่ครบและตัวกรอง"""
    social = HourCategory(name="กิจกรรมเพื่อสังคม", required_hours=10)
    sport = HourCategory(name="กิจกรรมกีฬา", required_hours=5)
    session.add_all([social, sport])
    session.commit()
    session.refresh(social)
    session.refresh(sport)

    sub = HourSubcategory(category_id=social.id, name="จิตอาสา", required_hours=10)
    session.add(sub)
    session.commit()
    session.refresh(sub)

    activity = Activity(
        name="ค่ายอาสา",
        activity_type="จิตอาสา",
        is_required=False,
        max_participants=100,
        start_at=datetime(2026, 6, 1, 9, 0),
        location="หอประชุม",
        hours=10,
        subcategory_id=sub.id,
        created_by=_staff_id(session),
        approval_status=ApprovalStatus.approved,
    )
    session.add(activity)
    session.commit()
    session.refresh(activity)

    a = Student(student_id="660001", full_name="เอ ครบหมวดสังคม", faculty="วิศวกรรมศาสตร์",
                major="คอมพิวเตอร์", year_level=4, status=StudentStatus.active)
    b = Student(student_id="660002", full_name="บี ยังไม่มีชั่วโมง", faculty="วิทยาศาสตร์",
                major="เคมี", year_level=1, status=StudentStatus.active)
    session.add_all([a, b])
    session.commit()
    session.refresh(a)
    session.refresh(b)

    session.add(
        Participation(
            student_id=a.id,
            activity_id=activity.id,
            evidence_status=EvidenceStatus.approved,
            hours_earned=10,
        )
    )
    session.commit()

    create_gold_layer(session)
    return {"a": a, "b": b, "social": social, "sport": sport}


# --------------------------- ข้อมูลในรายงาน ---------------------------
def test_rows_cover_every_student_and_category(session, world):
    rows = fetch_report_rows(session)

    # 2 นิสิต × 2 หมวด — หมวดที่ยังไม่ได้ทำต้องมีแถวด้วย (0 / required) ไม่ใช่หายไป
    assert len(rows) == 4
    assert {r.student_code for r in rows} == {"660001", "660002"}
    assert {r.category_name for r in rows} == {"กิจกรรมเพื่อสังคม", "กิจกรรมกีฬา"}


def test_row_carries_hours_and_status(session, world):
    rows = {(r.student_code, r.category_name): r for r in fetch_report_rows(session)}

    done = rows[("660001", "กิจกรรมเพื่อสังคม")]
    assert (done.earned_hours, done.required_hours) == (10, 10)
    assert done.completed is True
    assert done.status_text == STATUS_COMPLETE

    missing = rows[("660001", "กิจกรรมกีฬา")]
    assert (missing.earned_hours, missing.required_hours) == (0, 5)
    assert missing.status_text == STATUS_INCOMPLETE

    assert rows[("660001", "กิจกรรมเพื่อสังคม")].full_name == "เอ ครบหมวดสังคม"
    assert rows[("660002", "กิจกรรมกีฬา")].faculty == "วิทยาศาสตร์"
    assert rows[("660002", "กิจกรรมกีฬา")].year_level == 1


def test_rows_respect_filters(session, world):
    only_science = fetch_report_rows(session, faculty="วิทยาศาสตร์")
    assert {r.student_code for r in only_science} == {"660002"}

    year_four = fetch_report_rows(session, year_level=4)
    assert {r.student_code for r in year_four} == {"660001"}

    assert fetch_report_rows(session, faculty="วิทยาศาสตร์", year_level=4) == []


def test_filter_note_is_readable():
    assert describe_filters(None, None) == "ทุกคณะ ทุกชั้นปี"
    assert "วิทยาศาสตร์" in describe_filters("วิทยาศาสตร์", None)
    assert "ชั้นปีที่ 2" in describe_filters(None, 2)


def test_filename_is_thai_with_date():
    name = report_filename("xlsx", datetime(2026, 8, 5))
    assert name == "รายงานชั่วโมงกิจกรรม_2026-08-05.xlsx"


# --------------------------- ไฟล์ Excel ---------------------------
def _read_workbook(content: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(content)).active


def test_xlsx_has_thai_headers_and_one_row_per_record(session, world):
    content = build_xlsx(fetch_report_rows(session), filter_note="ทุกคณะ ทุกชั้นปี")
    sheet = _read_workbook(content)

    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    header_index = next(i for i, row in enumerate(values) if row[0] == "รหัสนิสิต")

    assert values[0][0] == "รายงานชั่วโมงกิจกรรมสะสมของนิสิต"
    assert values[header_index] == COLUMN_HEADERS
    assert len(values) - header_index - 1 == 4  # 2 นิสิต × 2 หมวด


def test_xlsx_hours_are_numbers_not_text(session, world):
    """ต้องเป็นตัวเลขจริง ไม่งั้นผู้ใช้ SUM/เรียงลำดับใน Excel ไม่ได้"""
    content = build_xlsx(fetch_report_rows(session), filter_note="ทุกคณะ ทุกชั้นปี")
    sheet = _read_workbook(content)

    rows = [[c.value for c in row] for row in sheet.iter_rows()]
    data_rows = [r for r in rows if r[0] in {"660001", "660002"}]
    assert data_rows, "ไม่พบแถวข้อมูลในไฟล์"
    for row in data_rows:
        assert isinstance(row[5], (int, float))  # ชั่วโมงที่ได้
        assert isinstance(row[6], (int, float))  # ชั่วโมงที่ต้องการ
        assert row[7] in {STATUS_COMPLETE, STATUS_INCOMPLETE}


def test_xlsx_endpoint_returns_a_downloadable_file(client, tokens, world):
    response = client.get("/reports/student-hours.xlsx", headers=_admin_headers(tokens))

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    # ชื่อไฟล์ไทยต้องเดินทางมาแบบ percent-encoded (RFC 5987)
    assert "filename*=UTF-8''" in disposition
    assert "%E0%B8%A3%E0%B8%B2%E0%B8%A2%E0%B8%87%E0%B8%B2%E0%B8%99" in disposition  # "รายงาน"
    assert _read_workbook(response.content).max_row > 1


def test_xlsx_endpoint_applies_filters(client, tokens, world):
    response = client.get(
        "/reports/student-hours.xlsx",
        params={"faculty": "วิทยาศาสตร์"},
        headers=_admin_headers(tokens),
    )

    sheet = _read_workbook(response.content)
    codes = {row[0].value for row in sheet.iter_rows(min_col=1, max_col=1) if row[0].value}
    assert "660002" in codes
    assert "660001" not in codes


# --------------------------- ไฟล์ PDF ---------------------------
def test_pdf_endpoint_returns_a_real_pdf(client, tokens, world):
    response = client.get("/reports/student-hours.pdf", headers=_admin_headers(tokens))

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF-")
    assert response.content.rstrip().endswith(b"%%EOF")
    # มีเนื้อหาจริง ไม่ใช่หน้าเปล่า
    assert len(response.content) > 2000


def test_pdf_embeds_a_thai_font(client, tokens, world):
    """ฟอนต์ไทยต้องถูกฝังมากับไฟล์ ไม่งั้นเครื่องที่เปิดจะเห็นเป็นกล่องสี่เหลี่ยม"""
    content = client.get("/reports/student-hours.pdf", headers=_admin_headers(tokens)).content
    assert b"FontFile2" in content  # TrueType ที่ฝังมา


def test_pdf_reports_a_clear_error_when_no_thai_font(client, tokens, world, monkeypatch):
    """เครื่องที่ไม่มีฟอนต์ไทยต้องได้ error ที่บอกทางแก้ ไม่ใช่ PDF ที่อ่านไม่ออก"""
    from app import reports as reports_module
    from reportlab.pdfbase import pdfmetrics

    monkeypatch.setattr(reports_module, "find_thai_font", lambda: None)
    # ฟอนต์ที่เคยลงทะเบียนไว้จากเทสต์ก่อนหน้าอยู่ในสถานะ global ของ reportlab
    monkeypatch.setattr(pdfmetrics, "getRegisteredFontNames", lambda: [])

    response = client.get("/reports/student-hours.pdf", headers=_admin_headers(tokens))

    assert response.status_code == 503
    assert "ฟอนต์" in response.json()["detail"]


# --------------------------- สิทธิ์ ---------------------------
def test_staff_can_download_both_formats(client, world):
    # client ตัวปกติล็อกอินเป็น staff
    assert client.get("/reports/student-hours.xlsx").status_code == 200
    assert client.get("/reports/student-hours.pdf").status_code == 200


def test_students_cannot_download_reports(client, tokens, world):
    headers = _student_headers(tokens)
    assert client.get("/reports/student-hours.xlsx", headers=headers).status_code == 403
    assert client.get("/reports/student-hours.pdf", headers=headers).status_code == 403


def test_reports_require_authentication(client, world):
    no_auth = {"Authorization": ""}
    assert client.get("/reports/student-hours.xlsx", headers=no_auth).status_code == 401
